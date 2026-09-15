"""Importa el rol de turnos RRHH (Excel D.L. 276 / distintos regímenes)."""

from __future__ import annotations

import re
from datetime import datetime, time
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from . import database as db


def _hhmm(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return f"{value.hour:02d}:{value.minute:02d}"
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    text = str(value).strip()
    if not text or text.lower() in ("none", "nan"):
        return None
    match = re.search(r"(\d{1,2}):(\d{2})", text)
    if not match:
        return None
    return f"{int(match.group(1)):02d}:{int(match.group(2)):02d}"


def _kind(entry: str, exit_time: str, lunch_start: str | None, lunch_end: str | None) -> str:
    to_min = lambda hhmm: int(hhmm[:2]) * 60 + int(hhmm[3:5])
    if to_min(exit_time) < to_min(entry):
        return "overnight"
    if not lunch_start or not lunch_end or lunch_start == lunch_end:
        return "block"
    return "split"


def _shift(entry, exit_time, lunch_start=None, lunch_end=None, label="") -> dict | None:
    if not entry or not exit_time:
        return None
    kind = _kind(entry, exit_time, lunch_start, lunch_end)
    if kind != "split":
        lunch_start = None
        lunch_end = None
    return {
        "label": label or f"{entry}–{exit_time}",
        "entry_time": entry,
        "exit_time": exit_time,
        "lunch_start": lunch_start or "",
        "lunch_end": lunch_end or "",
        "schedule_kind": kind,
    }


def parse_rol_turnos(file_bytes: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    people: list[dict] = []
    for row in rows:
        if not row or len(row) < 5:
            continue
        name = str(row[1] or "").strip()
        dni_raw = str(row[2] or "").strip()
        dni = re.sub(r"\D", "", dni_raw)
        times = [_hhmm(cell) for cell in row[3:9]]
        times = [t for t in times if t]
        note = str(row[7] or row[8] or "").strip() if len(row) > 7 else ""
        if not name or not dni or len(dni) < 6 or not times:
            continue
        if len(dni) > 8 and not name.replace(" ", "").isalpha():
            continue
        shifts: list[dict] = []
        if len(times) >= 6:
            for i, label in enumerate(("Turno 1", "Turno 2", "Turno 3")):
                item = _shift(times[i * 2], times[i * 2 + 1], label=label)
                if item:
                    shifts.append(item)
        elif len(times) == 4:
            t1, t2, t3, t4 = times
            a, b = int(t1[:2]) * 60 + int(t1[3:]), int(t2[:2]) * 60 + int(t2[3:])
            c, d = int(t3[:2]) * 60 + int(t3[3:]), int(t4[:2]) * 60 + int(t4[3:])
            gap = (c - b) if c > b else 0
            looks_two_blocks = (b - a) >= 6 * 60 or (d - c) >= 6 * 60 or abs(b - c) <= 30 or gap >= 180
            if looks_two_blocks and not (c > b and (c - b) <= 150):
                s1 = _shift(t1, t2, label="Turno A")
                s2 = _shift(t3, t4, label="Turno B")
                shifts = [s for s in (s1, s2) if s]
            else:
                item = _shift(t1, t4, t2, t3, label="Jornada partida")
                if item:
                    shifts.append(item)
        else:
            item = _shift(times[0], times[1], label="Jornada continua")
            if item:
                shifts.append(item)
        if not shifts:
            continue
        if (
            len(shifts) == 1
            and shifts[0].get("entry_time") == "08:00"
            and shifts[0].get("lunch_start") == "13:00"
            and shifts[0].get("lunch_end") == "14:30"
            and shifts[0].get("exit_time") == "17:30"
        ):
            continue
        people.append({
            "name": name,
            "dni": dni.lstrip("0") or dni,
            "dni_raw": dni,
            "shifts": shifts,
            "notes": note,
        })
    return people


def _match_user(person: dict, users: list[dict]) -> dict | None:
    dni = person["dni"]
    dni_raw = person["dni_raw"]
    for user in users:
        uid = str(user.get("user_id") or "")
        uid_digits = re.sub(r"\D", "", uid).lstrip("0") or uid
        if uid == dni_raw or uid == dni or uid_digits == dni:
            return user
    name_tokens = set(re.sub(r"\s+", " ", person["name"].upper()).split())
    best, score = None, 0
    for user in users:
        tokens = set(re.sub(r"\s+", " ", str(user.get("name") or "").upper()).split())
        if not tokens:
            continue
        hit = len(name_tokens & tokens)
        if hit >= 2 and hit > score:
            best, score = user, hit
    return best if score >= 2 else None


def apply_rol_turnos(file_bytes: bytes) -> dict:
    people = parse_rol_turnos(file_bytes)
    users = db.get_users()
    applied = 0
    unmatched = []
    rotating = 0
    for person in people:
        user = _match_user(person, users)
        if not user:
            unmatched.append({"name": person["name"], "dni": person["dni_raw"]})
            continue
        primary = person["shifts"][0]
        extras = person["shifts"][1:]
        notes = person.get("notes") or ""
        if extras:
            rotating += 1
            notes = (notes + " · Turnos rotativos").strip(" ·")
        db.save_employee_schedule(
            user["user_id"],
            user.get("name") or person["name"],
            entry_time=primary["entry_time"],
            exit_time=primary["exit_time"],
            lunch_start=primary.get("lunch_start") or "",
            lunch_end=primary.get("lunch_end") or "",
            notes=notes or None,
            schedule_kind=primary.get("schedule_kind") or "split",
            extra_shifts=extras,
        )
        applied += 1
    return {
        "parsed": len(people),
        "applied": applied,
        "rotating": rotating,
        "unmatched": unmatched[:40],
        "unmatched_count": len(unmatched),
    }


def default_rol_path() -> Path:
    return Path(__file__).resolve().parent.parent / "recursos" / "ROL DE TURNOS D.L 276 EMPLEADOS.xlsx"
